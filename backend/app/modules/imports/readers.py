import csv
import io
import re
import struct
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import openpyxl
import pandas as pd
import xlrd
from charset_normalizer import from_bytes
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.compdoc import CompDoc, CompDocError

from app.core.config import get_settings
from app.core.errors import AppError
from app.modules.imports.validation import invalid

OLE = bytes.fromhex("D0CF11E0A1B11AE1")


def unsupported() -> AppError:
    return AppError(415, "IMPORT_FILE_UNSUPPORTED", "CSV, XLSX, XLS 파일만 올릴 수 있습니다.")


def encrypted() -> AppError:
    return AppError(
        422,
        "IMPORT_FILE_ENCRYPTED",
        "암호가 걸린 파일입니다. 카드사 사이트에서 암호 없이 내려받아 주세요.",
    )


def too_large() -> AppError:
    return AppError(
        413,
        "IMPORT_FILE_TOO_LARGE",
        f"파일은 {get_settings().import_max_file_mb}MB 이하만 올릴 수 있습니다.",
    )


def cell_string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = (
            value.strftime("%Y-%m-%d %H:%M:%S")
            if value.time().isoformat() != "00:00:00"
            else value.strftime("%Y-%m-%d")
        )
    elif isinstance(value, date):
        value = value.isoformat()
    elif isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if not number.is_finite():
            return ""
        value = format(number, "f")
        if "." in value:
            value = value.rstrip("0").rstrip(".")
    return re.sub(r"\s+", " ", str(value)).strip()


def decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    match = from_bytes(data).best()
    if match is None:
        raise unsupported()
    return str(match).lstrip("\ufeff")


@dataclass
class Workbook:
    file_type: str
    sheets: dict[str, list[list[str]]]


def read_file(data: bytes) -> Workbook:
    if len(data) > get_settings().import_max_file_mb * 1024 * 1024:
        raise too_large()
    if not data:
        raise unsupported()
    try:
        if data.startswith(b"PK\x03\x04"):
            with ZipFile(io.BytesIO(data)) as archive:
                if "[Content_Types].xml" not in archive.namelist():
                    raise unsupported()
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    invalid("file", "압축을 푼 표가 너무 큽니다.")
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                return Workbook(
                    "xlsx",
                    {
                        sheet.title: [
                            [cell_string(v) for v in row]
                            for row in sheet.iter_rows(values_only=True)
                        ]
                        for sheet in book.worksheets
                    },
                )
            finally:
                book.close()
        if data.startswith(OLE):
            # OLE 디렉터리의 UTF-16 스트림 이름도 확인해 손상된 암호 파일을 안내한다.
            if "EncryptedPackage".encode("utf-16-le") in data:
                raise encrypted()
            compound = CompDoc(data, logfile=io.StringIO())
            if compound.locate_named_stream("EncryptedPackage")[0] is not None:
                raise encrypted()
            book = xlrd.open_workbook(file_contents=data, logfile=io.StringIO())
            try:
                sheets = {}
                for sheet in book.sheets():
                    rows = []
                    for row in sheet.get_rows():
                        rows.append(
                            [
                                cell_string(
                                    xlrd.xldate_as_datetime(c.value, book.datemode)
                                    if c.ctype == xlrd.XL_CELL_DATE
                                    else c.value
                                )
                                for c in row
                            ]
                        )
                    sheets[sheet.name] = rows
                return Workbook("xls", sheets)
            finally:
                book.release_resources()
        if data.startswith((b"%PDF", b"\x89PNG", b"\xff\xd8", b"GIF8")):
            raise unsupported()
        content = decode(data)
        if content.lstrip().startswith("<") and re.search(r"<table\b", content, re.I):
            # 헤더도 원본 행으로 유지한다. lxml이 표를 읽고 pandas가 셀 값으로 바꾼다.
            from lxml import etree, html

            root = html.fromstring(content, parser=html.HTMLParser(no_network=True))
            tables = root.xpath("//table[not(ancestor::table)]")
            sheets = {}
            for index, table in enumerate(tables, 1):
                for th in table.xpath(".//th"):
                    th.tag = "td"
                for thead in table.xpath(".//thead"):
                    thead.tag = "tbody"
                frame = pd.read_html(
                    io.StringIO(etree.tostring(table, encoding="unicode")),
                    flavor="lxml",
                    header=None,
                    keep_default_na=False,
                )[0]
                sheets[f"Sheet{index}"] = [[cell_string(v) for v in row] for row in frame.values]
            return Workbook("html", sheets)
        if any(ord(c) < 32 and c not in "\r\n\t" for c in content):
            raise unsupported()
        sample = "\n".join(content.splitlines()[:5])
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
        except csv.Error:
            scores = {
                d: sum(line.count(d) for line in sample.splitlines()) for d in (",", "\t", ";")
            }
            delimiter = max(scores, key=scores.get)
        csv.field_size_limit(
            max(csv.field_size_limit(), get_settings().import_max_file_mb * 1024 * 1024)
        )
        rows = [
            [cell_string(v) for v in row]
            for row in csv.reader(io.StringIO(content), delimiter=delimiter)
        ]
        return Workbook("csv", {"Sheet1": rows})
    except (
        BadZipFile,
        ValueError,
        KeyError,
        OSError,
        xlrd.XLRDError,
        CompDocError,
        csv.Error,
        ParseError,
        InvalidFileException,
        struct.error,
    ) as exc:
        if "encrypt" in str(exc).lower():
            raise encrypted() from exc
        raise unsupported() from exc
